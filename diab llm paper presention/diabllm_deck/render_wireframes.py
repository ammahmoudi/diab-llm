from pathlib import Path
from io import BytesIO
import math
import textwrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PPTX = Path(__file__).with_name("DiabLLM_paper_presentation_Amirhossein_Mahmoudi.pptx")
OUT = Path(__file__).with_name("qa_previews")
OUT.mkdir(exist_ok=True)

prs = Presentation(PPTX)
SW, SH = prs.slide_width, prs.slide_height
WIDTH, HEIGHT = 1280, 720
sx, sy = WIDTH / SW, HEIGHT / SH


def font(size, bold=False):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, max(7, int(size)))
    return ImageFont.load_default()


def color(rgb, fallback=(255, 255, 255)):
    try:
        if rgb is not None:
            value = str(rgb)
            if len(value) == 6:
                return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))
    except Exception:
        pass
    return fallback


def shape_fill(shape):
    try:
        return color(shape.fill.fore_color.rgb, (247, 250, 252))
    except Exception:
        return (247, 250, 252)


def shape_line(shape):
    try:
        return color(shape.line.color.rgb, (190, 205, 214))
    except Exception:
        return (190, 205, 214)


def render_text(draw, shape, box):
    text = shape.text.strip()
    if not text:
        return
    x1, y1, x2, y2 = box
    try:
        paragraph = shape.text_frame.paragraphs[0]
        run = paragraph.runs[0] if paragraph.runs else None
        point_size = float(run.font.size.pt) if run and run.font.size else 12
        bold = bool(run.font.bold) if run else False
        fill = color(run.font.color.rgb, (20, 33, 61)) if run else (20, 33, 61)
    except Exception:
        point_size, bold, fill = 12, False, (20, 33, 61)
    px_size = max(7, point_size * HEIGHT / 540)
    fnt = font(px_size, bold)
    max_chars = max(8, int((x2 - x1) / max(5, px_size * 0.55)))
    wrapped = []
    for line in text.splitlines() or [text]:
        wrapped.extend(textwrap.wrap(line, width=max_chars) or [""])
    line_h = max(8, int(px_size * 1.15))
    max_lines = max(1, int((y2 - y1) / line_h))
    clipped = wrapped[:max_lines]
    if len(wrapped) > max_lines and clipped:
        clipped[-1] = clipped[-1][: max(1, max_chars - 1)] + "..."
    draw.multiline_text((x1 + 4, y1 + 2), "\n".join(clipped), font=fnt, fill=fill, spacing=1)


def render_slide(slide, index):
    bg = (247, 250, 252)
    try:
        bg = color(slide.background.fill.fore_color.rgb, bg)
    except Exception:
        pass
    canvas = Image.new("RGB", (WIDTH, HEIGHT), bg)
    draw = ImageDraw.Draw(canvas)
    for shape in slide.shapes:
        x1 = int(shape.left * sx)
        y1 = int(shape.top * sy)
        x2 = int((shape.left + shape.width) * sx)
        y2 = int((shape.top + shape.height) * sy)
        if x2 <= 0 or y2 <= 0 or x1 >= WIDTH or y1 >= HEIGHT:
            continue
        box = (x1, y1, x2, y2)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                image = Image.open(BytesIO(shape.image.blob)).convert("RGB")
                image.thumbnail((max(1, x2 - x1), max(1, y2 - y1)), Image.Resampling.LANCZOS)
                px = x1 + (x2 - x1 - image.width) // 2
                py = y1 + (y2 - y1 - image.height) // 2
                canvas.paste(image, (px, py))
            except Exception:
                draw.rectangle(box, fill=(225, 233, 238), outline=(120, 140, 150), width=1)
        elif shape.shape_type == MSO_SHAPE_TYPE.CHART:
            draw.rounded_rectangle(box, radius=8, fill=(255, 255, 255), outline=(190, 205, 214), width=1)
            draw.line((x1 + 45, y2 - 35, x2 - 20, y2 - 35), fill=(130, 145, 155), width=1)
            draw.line((x1 + 45, y1 + 20, x1 + 45, y2 - 35), fill=(130, 145, 155), width=1)
            draw.text((x1 + 55, y1 + 24), "Native editable chart", font=font(12, True), fill=(0, 124, 131))
        else:
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                draw.rounded_rectangle(box, radius=6, fill=shape_fill(shape), outline=shape_line(shape), width=1)
            elif shape.shape_type == MSO_SHAPE_TYPE.LINE:
                draw.line((x1, y1, x2, y2), fill=shape_line(shape), width=2)
            if hasattr(shape, "text"):
                render_text(draw, shape, box)
    draw.text((WIDTH - 42, HEIGHT - 22), str(index), font=font(10, True), fill=(100, 116, 139))
    target = OUT / f"slide-{index:02d}.png"
    canvas.save(target)
    return canvas


slides = [render_slide(slide, i) for i, slide in enumerate(prs.slides, 1)]
thumb_w, thumb_h = 384, 216
cols = 3
rows = math.ceil(len(slides) / cols)
contact = Image.new("RGB", (cols * thumb_w, rows * (thumb_h + 28)), (226, 232, 236))
draw = ImageDraw.Draw(contact)
for i, image in enumerate(slides):
    thumb = image.resize((thumb_w, thumb_h), Image.Resampling.LANCZOS)
    x = (i % cols) * thumb_w
    y = (i // cols) * (thumb_h + 28)
    contact.paste(thumb, (x, y))
    draw.text((x + 8, y + thumb_h + 5), f"Slide {i + 1}", font=font(12, True), fill=(20, 33, 61))
contact.save(Path(__file__).with_name("DiabLLM_contact_sheet.png"))
print(f"Rendered {len(slides)} slide wireframes and contact sheet")
